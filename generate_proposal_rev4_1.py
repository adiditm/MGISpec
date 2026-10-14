import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('DOKUMEN SPESIFIKASI & ACUAN FITUR\nPEMBUATAN WEBSITE KATALOG PRODUK (REV 4.1)')
run.bold = True
run.font.size = Pt(14)

doc.add_paragraph('Kepada Yth.\nBpk Bram Setiawan Liem\nDi Tempat')
p_sub = doc.add_paragraph()
p_sub.add_run('Perihal: Dokumen Acuan Fitur & Ruang Lingkup Sistem Website Katalog MGISpec').bold = True
doc.add_paragraph('---')

# Section 1
p1_title = doc.add_paragraph()
p1_title.add_run('1. Latar Belakang & Tujuan').bold = True
doc.add_paragraph('Dokumen ini disusun sebagai acuan teknis dan spesifikasi fitur (non-pricing reference) dalam pembuatan sistem informasi katalog produk kacamata medis MGISpec berbasis Web Responsive. Sistem ini dirancang untuk memudahkan manajemen katalog produk secara terpusat melalui Content Management System (CMS) serta menyajikan antarmuka modern dua bahasa untuk target pasar Indonesia dan Filipina.')

# Section 2
p2_title = doc.add_paragraph()
p2_title.add_run('2. Ruang Lingkup Pekerjaan (Scope of Work)').bold = True
doc.add_paragraph('Desain Web Responsive: Tampilan optimal di PC, tablet, maupun smartphone.', style='List Bullet')
doc.add_paragraph('Fitur Dual Bahasa (Bilingual): Dukungan bahasa Indonesia (ID) dan Filipina (PH) untuk teks statis maupun data dinamis.', style='List Bullet')
doc.add_paragraph('Content Management System (CMS): Panel admin untuk mengelola banner slider, program promo/keunggulan, serta katalog produk kacamata.', style='List Bullet')
doc.add_paragraph('Fokus Katalog & Profil: Sistem tidak menggunakan shopping cart transaksi langsung, melainkan diarahkan pada informasi spesifikasi dan pemesanan via WhatsApp.', style='List Bullet')
doc.add_paragraph('Struktur 2 Halaman Utama:', style='List Bullet')

p2_1 = doc.add_paragraph('1. HOME (BERANDA): Menampilkan header dengan pemilih bahasa & link login, hero slider layar penuh (full-screen responsive), kartu manfaat/promo, dan carousel thumbnail produk tanpa harga.')
p2_1.paragraph_format.left_indent = Inches(0.4)
p2_2 = doc.add_paragraph('2. KATALOG (DETAIL PRODUK): Menampilkan galeri foto interaktif, ulasan teknologi lensa (ion/anti radiasi), dimensi frame, dan tombol konsultasi/pemesanan langsung ke WhatsApp.')
p2_2.paragraph_format.left_indent = Inches(0.4)

# Section 3
p3_title = doc.add_paragraph()
p3_title.add_run('3. Rincian Fitur & Modul Sistem').bold = True

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'No'
hdr_cells[1].text = 'Modul Sistem'
hdr_cells[2].text = 'Rincian Fitur & Deskripsi'
hdr_cells[0].paragraphs[0].runs[0].bold = True
hdr_cells[1].paragraphs[0].runs[0].bold = True
hdr_cells[2].paragraphs[0].runs[0].bold = True

modules = [
    ("1", "Front-End Beranda (Home)", "• Header: Logo resmi HD, Language Switcher (ID/PH), Link Login Admin\n• Fullscreen Hero Slider: Tampilan responsif desktop & smartphone\n• Section Manfaat: Kartu keunggulan medis & sertifikasi MDA/FDA\n• Section Katalog Di Atas Footer: Thumbnail seluruh produk tanpa harga\n• Footer: Profil resmi, kontak, dan copyright resmi MGISpec"),
    ("2", "Front-End Detail Produk (Catalog)", "• Galeri Foto Produk: Multi-thumbnail interaktif\n• Informasi Spesifikasi: Material frame, bridge, bobot, panjang tangkai\n• Edukasi Teknologi: Lensa oksigen, anti radiasi blue light & UV400\n• Call-to-Action (CTA): Tombol langsung pemesanan via WhatsApp"),
    ("3", "Content Management System (Admin)", "• Autentikasi Admin: Login dashboard yang aman\n• Manajemen Slider: Upload gambar terpisah desktop/mobile & teks link\n• Manajemen Produk: Tambah, edit, dan hapus foto serta spesifikasi\n• Pengaturan Web: Nomor WhatsApp CS, email, dan link media sosial"),
    ("4", "Sistem Multibahasa (Bilingual ID/PH)", "• Deteksi Bahasa: Pilihan bahasa via URL atau selector bendera\n• Penerjemahan Dinamis: Form input di CMS mendukung bahasa ID & PH\n• Penerjemahan Antarmuka: Translasi otomatis teks tombol dan navigasi")
]

for no, mod, desc in modules:
    row_cells = table.add_row().cells
    row_cells[0].text = no
    row_cells[1].text = mod
    row_cells[1].paragraphs[0].runs[0].bold = True
    row_cells[2].text = desc

doc.add_paragraph('') # Spacing

# Section 4
p4_title = doc.add_paragraph()
p4_title.add_run('4. Estimasi Waktu Pengerjaan (Timeline)').bold = True
doc.add_paragraph('Total waktu pengerjaan diperkirakan memakan waktu 1 - 2 Minggu, dengan tahapan:')
doc.add_paragraph('1. Minggu 1: Setup arsitektur backend, database migration, dan pembangunan panel CMS Admin.', style='List Number')
doc.add_paragraph('2. Minggu 2: Integrasi front-end responsif, sistem dual bahasa, input materi katalog, dan testing di hosting Dewaweb.', style='List Number')

doc.add_paragraph('\nDemikian dokumen spesifikasi acuan fitur ini kami susun sebagai pedoman teknis pengembangan website katalog MGISpec.\n\nTerima kasih atas perhatian dan kerja samanya.')
doc.add_paragraph('\nHormat kami,\n\n\n\n(Andreas Didit M.)')
doc.add_page_break()

# Attachment
p_att = doc.add_paragraph()
p_att.add_run('Lampiran: Contoh Desain Mockup').bold = True
p_att.runs[0].font.size = Pt(14)
p_att.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('1. Tampilan Halaman Depan (Banner & Produk Unggulan)').bold = True
doc.add_picture('/Users/didit/.gemini/antigravity/brain/ab946cb2-4f8c-4f07-b4aa-be5903d85fe0/mockup_home.png', width=Inches(6.0))

doc.add_page_break()

doc.add_paragraph('2. Tampilan Halaman Detail Katalog (Spesifikasi Produk)').bold = True
doc.add_picture('/Users/didit/.gemini/antigravity/brain/ab946cb2-4f8c-4f07-b4aa-be5903d85fe0/mockup_details.png', width=Inches(6.0))

doc.save('Penawaran_Katalog_Kacamata_Rev4.1.docx')
print("Successfully generated Penawaran_Katalog_Kacamata_Rev4.1.docx")
