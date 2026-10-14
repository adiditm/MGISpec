import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('PROPOSAL PENAWARAN PEMBUATAN WEBSITE KATALOG PRODUK')
run.bold = True
run.font.size = Pt(14)

doc.add_paragraph('Kepada Yth.\nBpk Bram Setiawan Liem\nDi Tempat')
doc.add_paragraph('Perihal: Penawaran Jasa Pembuatan Website Katalog Produk Kacamata')
doc.add_paragraph('---')

# Section 1
p1_title = doc.add_paragraph()
p1_title.add_run('1. Latar Belakang & Tujuan').bold = True
doc.add_paragraph('Proposal ini disusun untuk menawarkan solusi pembuatan sistem informasi dan katalog produk kacamata berbasis Web (Responsive). Sistem ini dirancang untuk memudahkan pengelolaan produk secara terpusat melalui Content Management System (CMS) yang sederhana dan efisien.')

# Section 2
p2_title = doc.add_paragraph()
p2_title.add_run('2. Ruang Lingkup Pekerjaan (Scope of Work)').bold = True
p2_sub = doc.add_paragraph()
p2_sub.add_run('A. Website & CMS (Backend)').bold = True

doc.add_paragraph('Desain Web Responsive: Tampilan optimal di PC, tablet, maupun smartphone.', style='List Bullet')
doc.add_paragraph('Fitur Dual Bahasa (Bilingual): Dukungan bahasa Indonesia dan Filipina (Tagalog/English PH) untuk ekspansi target pasar yang lebih luas.', style='List Bullet')
doc.add_paragraph('Content Management System (CMS): Panel admin untuk mengelola (Upload/Edit/Hapus) data katalog produk dan gambar di halaman depan.', style='List Bullet')
doc.add_paragraph('Struktur Halaman:', style='List Bullet')

p2_1 = doc.add_paragraph('1. HOME: Halaman depan menampilkan produk unggulan (thumbnail) dan banner.')
p2_1.paragraph_format.left_indent = Inches(0.5)
p2_2 = doc.add_paragraph('2. KATALOG (PRODUCT DETAILS): Halaman detail produk yang dinamis menampilkan informasi produk, gambar, harga, dan deskripsi/spesifikasi.')
p2_2.paragraph_format.left_indent = Inches(0.5)

doc.add_paragraph('(Catatan: Pekerjaan kali ini difokuskan hanya untuk 2 halaman utama tersebut beserta sistem CMS-nya).').italic = True

# Section 3
p3_title = doc.add_paragraph()
p3_title.add_run('3. Rencana Anggaran Biaya').bold = True

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'No'
hdr_cells[1].text = 'Deskripsi Pekerjaan'
hdr_cells[2].text = 'Biaya (Rp)'

row_cells = table.add_row().cells
row_cells[0].text = '1'
row_cells[1].text = 'Pembuatan Website (Home & Katalog) & CMS (Admin)\n- Responsive web design\n- Fitur Dual Bahasa (ID & PH)\n- Pengelolaan produk dinamis (CMS)'
row_cells[2].text = '5.000.000'

total_row = table.add_row().cells
total_row[1].text = 'TOTAL BIAYA'
total_row[1].paragraphs[0].runs[0].bold = True
total_row[2].text = '5.000.000'
total_row[2].paragraphs[0].runs[0].bold = True

doc.add_paragraph('') # Spacing

# Section 4
p4_title = doc.add_paragraph()
p4_title.add_run('4. Estimasi Waktu Pengerjaan (Timeline)').bold = True
doc.add_paragraph('Total waktu pengerjaan diperkirakan memakan waktu 1 - 2 Minggu, dengan rincian:')
doc.add_paragraph('1. Minggu 1: Pengumpulan materi, perancangan desain awal, dan pembuatan antarmuka (Front-end).', style='List Number')
doc.add_paragraph('2. Minggu 2: Pengembangan CMS, integrasi database, dan testing.', style='List Number')

# Section 5
p5_title = doc.add_paragraph()
p5_title.add_run('5. Tahapan Pembayaran').bold = True
doc.add_paragraph('Tahap I (DP 50%): Rp 2.500.000,- (Saat kontrak disetujui & pekerjaan dimulai).', style='List Bullet')
doc.add_paragraph('Tahap II (50%): Rp 2.500.000,- (Saat Website dan CMS selesai dikembangkan dan siap di-testing/live).', style='List Bullet')

doc.add_paragraph('\nDemikian penawaran ini kami sampaikan. Jika ada pertanyaan lebih lanjut atau ingin menyesuaikan anggaran, silakan menghubungi kami.\n\nTerima kasih atas perhatian dan kerja samanya.')
doc.add_paragraph('\nHormat kami,\n\n\n\n(Andreas Didit M.)')
doc.add_page_break()

# Attachment
p_att = doc.add_paragraph()
p_att.add_run('Lampiran: Contoh Desain Mockup').bold = True
p_att.runs[0].font.size = Pt(14)
p_att.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('1. Tampilan Halaman Depan (Banner & Produk Unggulan)').bold = True
doc.add_picture('/Users/didit/.gemini/antigravity/brain/ab946cb2-4f8c-4f07-b4aa-be5903d85fe0/mockup_home.png', width=Inches(6.0))

# Insert Page Break to fix formatting issue
doc.add_page_break()

doc.add_paragraph('2. Tampilan Halaman Detail Katalog (Spesifikasi Produk)').bold = True
doc.add_picture('/Users/didit/.gemini/antigravity/brain/ab946cb2-4f8c-4f07-b4aa-be5903d85fe0/mockup_details.png', width=Inches(6.0))

doc.save('Penawaran_Katalog_Kacamata_Rev6.docx')
